"""첨부된 서울 공공데이터를 앱용 표준 CSV.gz로 변환한다."""

from __future__ import annotations

import argparse
import csv
import gzip
import re
from collections.abc import Iterable
from pathlib import Path

import openpyxl

DEFAULT_DOWNLOADS = Path.home() / "Downloads"


def _write_rows(path: Path, fieldnames: list[str], rows: Iterable[dict[str, object]]) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    count = 0
    with gzip.open(path, "wt", encoding="utf-8", newline="") as target:
        writer = csv.DictWriter(target, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)
            count += 1
    return count


def _number(value: object) -> float | None:
    try:
        return float(value) if value not in (None, "") else None
    except (TypeError, ValueError):
        return None


def prepare_trees(source: Path, target: Path) -> int:
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)

    def rows():
        for sheet in workbook:
            for row in sheet.iter_rows(min_row=4, values_only=True):
                longitude, latitude = _number(row[5]), _number(row[6])
                if longitude is None or latitude is None:
                    continue
                species = str(row[2] or "").strip()
                yield {
                    "district": str(row[0] or "").strip(),
                    "route": str(row[1] or "").strip(),
                    "species": species,
                    "is_female_ginkgo": str(species == "은행나무 암나무").lower(),
                    "address": str(row[3] or row[4] or "").strip(),
                    "longitude": longitude,
                    "latitude": latitude,
                    "is_sample": "false",
                }

    return _write_rows(
        target,
        [
            "district",
            "route",
            "species",
            "is_female_ginkgo",
            "address",
            "longitude",
            "latitude",
            "is_sample",
        ],
        rows(),
    )


def prepare_shades(source: Path, target: Path) -> int:
    sheet = openpyxl.load_workbook(source, read_only=True, data_only=True).active

    def rows():
        for row in sheet.iter_rows(min_row=9, values_only=True):
            longitude, latitude = _number(row[9]), _number(row[10])
            if longitude is None or latitude is None:
                continue
            yield {
                "id": str(row[0] or "").strip(),
                "type": str(row[1] or "").strip(),
                "district": str(row[3] or "").strip(),
                "name": str(row[6] or "").strip(),
                "address": str(row[7] or row[8] or "").strip(),
                "longitude": longitude,
                "latitude": latitude,
                "is_sample": "false",
            }

    return _write_rows(
        target,
        ["id", "type", "district", "name", "address", "longitude", "latitude", "is_sample"],
        rows(),
    )


def prepare_bikes(source: Path, target: Path) -> int:
    sheet = openpyxl.load_workbook(source, read_only=True, data_only=True).active

    def rows():
        for row in sheet.iter_rows(min_row=6, values_only=True):
            latitude, longitude = _number(row[4]), _number(row[5])
            if latitude is None or longitude is None or row[0] is None:
                continue
            capacity = sum(int(_number(value) or 0) for value in (row[7], row[8]))
            yield {
                "station_id": str(row[0]).strip(),
                "name": str(row[1] or "").strip(),
                "district": str(row[2] or "").strip(),
                "address": str(row[3] or "").strip(),
                "latitude": latitude,
                "longitude": longitude,
                "capacity": capacity,
                "availability_is_realtime": "false",
                "is_sample": "false",
            }

    return _write_rows(
        target,
        [
            "station_id",
            "name",
            "district",
            "address",
            "latitude",
            "longitude",
            "capacity",
            "availability_is_realtime",
            "is_sample",
        ],
        rows(),
    )


def prepare_streetlights(source: Path, target: Path) -> int:
    def rows():
        with source.open(encoding="cp949", newline="") as stream:
            for row in csv.DictReader(stream):
                latitude, longitude = _number(row.get("위도")), _number(row.get("경도"))
                if latitude is None or longitude is None:
                    continue
                yield {
                    "id": str(row.get("관리번호") or "").strip(),
                    "latitude": latitude,
                    "longitude": longitude,
                    "is_sample": "false",
                }

    return _write_rows(target, ["id", "latitude", "longitude", "is_sample"], rows())


def _road_name(location: str) -> str:
    prefix = location.split("(", 1)[0].strip()
    return re.sub(r"\s+", "", prefix)


def prepare_heating(source: Path, target: Path) -> int:
    def rows():
        with source.open(encoding="cp949", newline="") as stream:
            for row in csv.DictReader(stream):
                location = str(row.get("설치위치 (시점 ~ 종점)") or "").strip()
                yield {
                    "id": str(row.get("연번") or "").strip(),
                    "district": str(row.get("관리기관") or "").strip(),
                    "installed": str(row.get("설치연도") or "").strip(),
                    "location": location,
                    "road_name": _road_name(location),
                    "length_m": _number(row.get("설치연장(m)")) or 0,
                    "geometry_verified": "false",
                    "is_sample": "false",
                }

    return _write_rows(
        target,
        [
            "id",
            "district",
            "installed",
            "location",
            "road_name",
            "length_m",
            "geometry_verified",
            "is_sample",
        ],
        rows(),
    )


def prepare_safe_return(source: Path, target: Path) -> int:
    """안심귀갓길 통합 SHP의 시설 지점을 좌표 표준 형식으로 변환한다."""
    import geopandas as gpd

    frame = gpd.read_file(source).to_crs("EPSG:4326")

    def rows():
        for index, geometry in enumerate(frame.geometry):
            if geometry is None or geometry.is_empty or geometry.geom_type != "Point":
                continue
            yield {
                "id": f"safe-return-{index + 1}",
                "latitude": geometry.y,
                "longitude": geometry.x,
                "label": "안심귀갓길 연계 시설",
                "is_sample": "false",
            }

    return _write_rows(
        target, ["id", "latitude", "longitude", "label", "is_sample"], rows()
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source-dir", type=Path, default=DEFAULT_DOWNLOADS)
    parser.add_argument("--output-dir", type=Path, default=Path("data/processed"))
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    jobs = {
        "trees.csv.gz": prepare_trees(
            args.source_dir / "2026 서울시 가로수 위치정보.xlsx", args.output_dir / "trees.csv.gz"
        ),
        "shades.csv.gz": prepare_shades(
            args.source_dir / "폭염저감시설 목록_그늘막(2026.7.31.).xlsx",
            args.output_dir / "shades.csv.gz",
        ),
        "bike_stations.csv.gz": prepare_bikes(
            args.source_dir / "공공자전거 대여소 정보(26.6월 기준).xlsx",
            args.output_dir / "bike_stations.csv.gz",
        ),
        "streetlights.csv.gz": prepare_streetlights(
            args.source_dir / "서울시 가로등 위치 정보.csv",
            args.output_dir / "streetlights.csv.gz",
        ),
        "heating.csv.gz": prepare_heating(
            args.source_dir / "자치구별 도로열선 설치현황_20260531.csv",
            args.output_dir / "heating.csv.gz",
        ),
        "safe_return.csv.gz": prepare_safe_return(
            args.source_dir
            / "안심귀갓길 서비스 통합데이터_SHP"
            / "안심귀갓길 서비스 통합데이터.shp",
            args.output_dir / "safe_return.csv.gz",
        ),
    }
    for name, count in jobs.items():
        print(f"{name}: {count:,} rows")


if __name__ == "__main__":
    main()
